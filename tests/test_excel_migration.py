from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from backend.migration import ExcelImporter
from backend.models import MigraineEntry


HEADERS = [
    "Datum", "Auslöser (1-6)", "Stärke (0-10)", "Dauer (h)", "Typ", "Seite",
    "Vorboten (Codes)", "Erbrechen", "Übelkeit", "Lärmscheu", "Lichtscheu",
    "Geruchsempfindlich", "Andere Symptome (Codes)", "Medikament", "Aimovig Spritze",
    "Momeallerg Nasenspray", "Amitriptylin neuraxpharm", "Hat geholfen?", "Notizen",
]


def workbook_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kopfschmerzkalender"
    sheet.append([None] * len(HEADERS))
    sheet.append([None] * len(HEADERS))
    sheet.append([None] * len(HEADERS))
    sheet.append(HEADERS)
    sheet.append([date(2026, 6, 1), None, None, None, None, None, None, None, None, None, None, None, None, None, None, "x", None, None, None])
    exact_note = "Originalnotiz\nmit Zeilenumbruch und 17–18 °C."
    sheet.append([date(2026, 6, 2), "5", 7, 6.5, "Dumpf / drückend", "Rechts", None, None, "x", None, None, None, None, "Eletriptan (2x)", None, "x", None, "Nein", exact_note])
    table = Table(displayName="KopfschmerzEintraege", ref="A4:S6")
    sheet.add_table(table)
    workbook.create_sheet("Zusatzinformationen")
    workbook.save(path)


def test_excel_import_is_idempotent_and_preserves_notes(session, user, tmp_path) -> None:
    path = tmp_path / "tracker.xlsx"
    workbook_fixture(path)
    importer = ExcelImporter(session, user.id, today=date(2026, 6, 2))

    first = importer.import_workbook(path)
    session.commit()
    second = importer.import_workbook(path)
    session.commit()

    assert first.imported == 1
    assert first.daily_records_imported == 2
    assert first.rejected == 0
    assert second.imported == 0
    assert second.skipped == 2
    assert second.observed_daily_rows == 2
    entries = session.query(MigraineEntry).all()
    assert len(entries) == 1
    assert entries[0].notes == "Originalnotiz\nmit Zeilenumbruch und 17–18 °C."
    assert entries[0].medications[0].name == "Eletriptan"
    assert entries[0].medications[0].dose == "2x"
